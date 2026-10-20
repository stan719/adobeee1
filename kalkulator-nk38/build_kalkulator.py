"""Generuje kalkulator NK38 w Excelu (Kalkulator_NK38_v1.xlsx).

Excel jest "podkładką" aplikacji: liczy dokładnie to samo co model_nk38.py
i aplikacja, ale każdy krok da się prześledzić.

Układ:
  Dane         - jedyne miejsce, gdzie się coś wpisuje (żółte pola)
  Podsumowanie - 4 warianty (2000/2500 kg, z/bez słupków), opcje, analiza liczby
                 osób, zestawienie do oferty, uwagi do sprawdzenia
  Obliczenia   - materiały (dawny arkusz "NK38 ...") + montaż (dawny "<Inw>-NK38")
  Cennik       - wszystkie ceny i stałe

Użycie:  python3 build_kalkulator.py [plik.xlsx] [--przyklad waw|krk]
"""
import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")
FILL_RES = PatternFill("solid", fgColor="E2EFDA")
FILL_WARN = PatternFill("solid", fgColor="F8CBAD")
FILL_BEST = PatternFill("solid", fgColor="C6EFCE")
F_HEAD = Font(bold=True, color="FFFFFF", size=13)
F_BOLD = Font(bold=True)
F_NOTE = Font(color="7F7F7F", size=9)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PLN = '#,##0 "zł"'
PLN2 = '#,##0.00 "zł"'
EUR = '#,##0 "€"'
PCT = "0.0%"
NUM = "#,##0.##"

GLEBOKOSCI = [(175, 0, 2000), (180, 500, 2000), (185, 550, 2000), (190, 600, 2000),
              (195, 650, 2000), (200, 700, 2000), (210, 750, 2250), (220, 800, 2250),
              (225, 850, 2250)]
WYMIARY = [("szer1", "Szerokość 260 (S) / 520 (D)"), ("szer2", "Szerokość 270 (S) / 540 (D)"),
           ("dl500", "Długość zagłębienia 500-510 cm"), ("dl540", "Długość zagłębienia 540 cm"),
           ("dl550", "Długość zagłębienia 550 cm")]
WARIANTY = ["2000 kg", "2500 kg", "2000 kg + słupki", "2500 kg + słupki"]
WZORY = ["Stary (poprawiony)", "Uproszczony"]
OSOBY_ANALIZA = [2, 3, 4, 5, 6, 7]

UWAGI = [
    "Hotel liczony jako (dni pracy + 1) x osoby - do sprawdzenia (może zapas na przyjazd na odbiór).",
    "Podatek 14% od robocizny, diet, GDW i zysku = wartość efektywna; 'koszt ludzki' w starym arkuszu używał x1,19 - do sprawdzenia.",
    "Systemy N.P. liczone jak D (siłowniki, składowe, dźwig, 7 osobodni), cena 85% D, 2 MP - do potwierdzenia.",
    "Czas podróży: wzór 'Stary (poprawiony)' = stary dla < 800 km, dla >= 800 km uwzględnia liczbę wyjazdów. "
    "'Uproszczony' = 2 x km / 650 km dziennie na wyjazd.",
    "Rozładunek sprzętu to jednorazowy koszt przyjazdu i wyjazdu wózka (nie mnoży się przez liczbę TIR-ów).",
    "Zysk B liczony od tych samych kosztów co zysk A (dawna kolumna N, bez zmiany ceny dźwigu).",
]

CENNIK = [
    ("Platformy", None),
    ("c_cena_D", "Platforma podwójna NK38-500-170-530 (4 MP)", 32500, PLN, "NK38!E32"),
    ("c_doplata_NK37", "Dopłata za NK37 / system", 700, PLN, "NK38!F1"),
    ("c_wsp_S", "Cena S jako ułamek ceny D", 0.61, PCT, "NK38!F3"),
    ("c_wsp_NP", "Cena systemu N.P. jako ułamek ceny D", 0.85, PCT, "NK38!E27"),
    ("c_szer_260", "Szerokość S 260 cm", 1170, PLN, "NK38!E8"),
    ("c_szer_270", "Szerokość S 270 cm", 1560, PLN, "NK38!E9"),
    ("c_szer_520", "Szerokość D 520 cm", 1820, PLN, "NK38!E34"),
    ("c_szer_540", "Szerokość D 540 cm", 3110, PLN, "NK38!E35"),
    ("c_dl_500", "Długość zagłębienia 500-510 cm", 2340, PLN, "NK38!E21/E47"),
    ("c_dl_wsp", "Dopłata za 1 cm długości (ułamek ceny platformy)", 0.00127, "0.00000", "NK38!E22 (4 kg/cm z 3150 kg)"),
    ("c_q2500", "Dopłata Q=2500 kg (domyślna)", 0.07, PCT, "NK38!D90"),
    ("Hydraulika", None),
    ("c_agregat", "Agregat 11,5 l/min (do 6 platform)", 6500, PLN, "NK38!E56"),
    ("c_olej", "Olej 150 l / agregat", 1500, PLN, "NK38!D134"),
    ("c_hydr_drobne", "Elementy drobne / agregat", 200, PLN, "NK38!D135"),
    ("c_transp_agregatu", "Transport agregatu", 500, PLN, "NK38!F57"),
    ("c_transp_silownikow", "Transport siłowników / platformę", 50, PLN, "NK38!F57"),
    ("c_sil_1700", "Siłownik standard (zagł. 170)", 1080, PLN, "NK38!D130"),
    ("c_sil_2000", "Siłownik do zagł. 2000 (też 175 cm)", 1175, PLN, "NK38!D129"),
    ("c_sil_2250", "Siłownik do zagł. 2250", 1250, PLN, "NK38!D128"),
    ("c_magistrala", "Magistrala EXTRA / 6 m", "=0.2*11+0.4*11+2*14+6*11+6*5+6*10", PLN2, "NK38!D133"),
    ("Barierki, elektryka, administracja", None),
    ("c_barierka", "Barierka na stronę (Metazet 538 zł)", 600, PLN, "NK38!E66"),
    ("c_maskownica", "Maskownica / bariera tylna (Metazet 608 zł)", 700, PLN, "NK38!E67"),
    ("c_skrzynka_zasil", "Elektryka - skrzynka zasilająca / agregat", 1000, PLN, "NK38!D121"),
    ("c_admin", "Administracja, koszty biura / agregat", 1800, PLN, "NK38!D125"),
    ("Opcje oferty", None),
    ("c_mal_S", "Malowanie proszkowe RAL - system S", 6400, PLN, "NK38!E92"),
    ("c_mal_D", "Malowanie proszkowe RAL - system D / NP", 13500, PLN, "NK38!E93"),
    ("c_mal_bar", "Malowanie barierki na stronę", 200, PLN, "NK38!E94"),
    ("c_slupek", "Słupek e-ładowarki do 3,7 kW / MP", 1700, PLN, "NK38!E97, KAL!O55"),
    ("c_slupek22", "Słupek e-ładowarki do 22 kW / MP", 6300, PLN, "NK38!E98"),
    ("c_slupek_rbd", "Robocizna słupka (roboczodni / MP)", 0.37, NUM, "KAL!B49"),
    ("Montaż - pracochłonność", None),
    ("c_osd_S", "Osobodni montażu / system S", 5.5, NUM, "KAL!F16"),
    ("c_osd_D", "Osobodni montażu / system D", 7, NUM, "KAL!F16"),
    ("c_osd_NP", "Osobodni montażu / system NP", 7, NUM, "KAL!F17"),
    ("c_osd_magistrala", "Utrudnienia: osobodni / magistrala 6 m", 0.5, NUM, "KAL!B48 ('tonaż')"),
    ("c_osd_barierka", "Utrudnienia: osobodni / barierka lub maskownica", 0.5, NUM, "KAL!B48 ('tonaż')"),
    ("c_rozl_TIR", "Dni rozładunku / TIR", 0.5, NUM, "KAL!E55 (komentarz)"),
    ("c_dzwig", "Spider Crane / dzień", 2400, PLN, "KAL!I25"),
    ("c_dzwig_D", "Systemów D (i NP) na 1 dzień dźwigu", 5, NUM, "KAL!J25"),
    ("c_dzwig_wspS", "Współczynnik S względem D (dźwig)", 0.6, NUM, "KAL!J25"),
    ("c_dni_tydzien", "Dni montażu w tygodniu (wyjazd co tydzień)", 5, NUM, "B55 = J39/5"),
    ("Montaż - dojazdy", None),
    ("c_km_dodatek", "Montaż: km dodane do odległości (nocleg, sklep)", 50, NUM, "WAW: 280 -> 330"),
    ("c_km_krakow", "Kraków: km na pomiary i montaż", 20, NUM, "KRK: 20 / 20"),
    ("c_km_dzien", "Wzór uproszczony: km dziennie", 650, NUM, "propozycja"),
    ("c_spalanie", "Spalanie auta ekipy l/100 km (1 auto)", 12, NUM, "KAL!I19"),
    ("c_paliwo", "Cena paliwa zł/l", 9.4, PLN2, "KAL!I19"),
    ("c_pomiary_osoby", "Pomiary - liczba osób (1 wyjazd)", 1, NUM, "KAL!F53"),
    ("c_odbiory_osoby", "Odbiory - liczba osób (1 wyjazd)", 2, NUM, "KAL!F58"),
    ("Montaż - koszty ekipy", None),
    ("c_stawka_1", "Stawka dzienna - 1. osoba", 500, PLN, "KAL!I47"),
    ("c_stawka_2", "Stawka dzienna - 2. osoba", 400, PLN, "KAL!I48"),
    ("c_stawka_n", "Stawka dzienna - każda kolejna osoba", 350, PLN, "KAL!I49:I53"),
    ("c_dieta", "Dieta dzienna (za granicą 49 €)", 45, PLN, "KAL!J57"),
    ("c_hotel", "Hotel / doba / osoba", 130, PLN, "KAL!J58"),
    ("Narzuty", None),
    ("c_kob", "KOB (udział w cenie)", 0.005, PCT, "KAL!J35"),
    ("c_gdw", "GDW (udział w cenie)", 0.025, PCT, "KAL!J36"),
    ("c_podatek", "Podatek efektywny od robocizny, diet, GDW i zysku", 0.14, PCT, "KAL!J37"),
]

SKLADOWE = [
    ("Łańcuch 20B-1 (mb)", 45, 9.6, 9.6, "Mechanika"),
    ("Ogniwo złączne proste 20B-1", 8, 4, 4, "Mechanika"),
    ("Koło łańcuchowe z piastą 20B-1 Z10", 36, 2, 2, "Mechanika"),
    ("Koło łańcuchowe z piastą 20B-1 Z16", 90, 2, 2, "Mechanika"),
    ("Listwa zębata M4", 207.28, 2, 2, "Mechanika"),
    ("Koło zębate z piastą M4 Z20", 40.14, 2, 2, "Mechanika"),
    ("Tuleja PAF40x260P10", 3.3, 2, 2, "Mechanika"),
    ("Tuleja PAF 5025 P10", 10.8, 2, 2, "Mechanika"),
    ("Tuleja TBR40x44x50", 10.95, 2, 2, "Mechanika"),
    ("Podkładka PTFE 50,5x69x2", 3.81, 2, 2, "Mechanika"),
    ("Podkładka PTFE 42x66x1,5", 2.84, 4, 4, "Mechanika"),
    ("Tuleja PAF40x400P10", 4.4, 4, 4, "Mechanika"),
    ("Podkładki gumowe pod stopę siłownika i słup (kpl)", "=2*25.71+2*22.85", 1, 1, "Mechanika"),
    ("Śruby i elementy złączne - kpl S", 345, 1, 0, "Mechanika"),
    ("Śruby i elementy złączne - kpl D", 546, 0, 1, "Mechanika"),
    ("Tulejki podłogi", 2.98, 110, 208, "Mechanika"),
    ("Kotwa M16 + ampułka", "=3.09+4.99", 20, 20, "Mechanika"),
    ("Podkładki gumowe pod kotwy (kpl)", 5.4, 1, 1, "Mechanika"),
    ("Transport materiałów (bez MTZ)", 300, 1, 1, "Mechanika"),
    ("Ogranicznik koła", 40, 1, 1, "Mechanika"),
    ("Rurki gotowe na platformie - S", 172, 1, 0, "Hydraulika"),
    ("Rurki gotowe na platformie - D", 207, 0, 1, "Hydraulika"),
    ("Elementy złączne Argo Hytos + magistrala 5 mb + wąż", "=120+55+65", 1, 1, "Hydraulika"),
    ("Elektryka - skrzynka krosowa", "=50+40", 1, 1, "Elektryka"),
    ("Elektryka - przewody + rurki S", "=14+9+25+5+9+30", 1, 0, "Elektryka"),
    ("Elektryka - przewody + rurki D", "=14+18+25+7+18+30", 0, 1, "Elektryka"),
    ("Elektryka - sterownia (puszka+stac+przyc+notaus+szyld)", "=70+14+34+32+20", 1, 1, "Elektryka"),
]

PRZYKLADY = {
    "waw": dict(in_nazwa="ul. Kasprowicza 18 Warszawa", in_nr="20261014-SNK",
                in_lok="Poza Krakowem", in_km=280, in_D=15, in_agregaty=3,
                in_magistrala=4, in_barierki=30, in_maskownice=5, gl_D={200: 15},
                in_tir=3, in_tir_cena=4000, in_rozl_cena=5100, in_osoby=4,
                in_zysk=0.125, in_zysk2=0.10, in_q2500=0.06, in_data=date(2026, 10, 14),
                in_opis="H-195/202 | L-530-544-584 | P-380 | Hpoj. G-180. D-180 | 1A4+1A5+1A6 | Bok-30 Tył-5 | mgstrl-4"),
    "krk": dict(in_nazwa="ul. Dekerta Kraków", in_nr="20261015-SNK", in_lok="Kraków",
                in_S=12, in_D=27, in_agregaty=7, in_magistrala=32, in_barierki=58,
                in_maskownice=4, gl_S={225: 12}, gl_D={225: 27}, in_tir=6,
                in_tir_cena=1800, in_rozl_cena=3000, in_osoby=5, in_zysk=0.13,
                in_zysk2=0.11, in_q2500=0.07, in_data=date(2026, 10, 15),
                in_opis="H-225 | L-530 | P-392/352? | Hpoj. G-170(185). D-200(185) | 5A6+1A5+1A4 | Bok-58 Tył-4 | mgstrl- 32"),
}


def build(inputs=None):
    inputs = inputs or {}
    wb = Workbook()
    names = {}

    def absref(cell):
        col = "".join(c for c in cell if c.isalpha())
        return f"${col}${cell[len(col):]}"

    def name(n, ws, cell):
        names[n] = (ws, cell)
        wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!{absref(cell)}")

    def name_range(n, ws, rng):
        a, b = rng.split(":")
        wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!{absref(a)}:{absref(b)}")

    def header(ws, text, cols):
        ws["B2"] = text
        ws["B2"].font = F_HEAD
        for c in range(2, 2 + cols):
            ws.cell(row=2, column=c).fill = FILL_HEAD
        ws.row_dimensions[2].height = 24

    def section(ws, row, text, cols=4):
        ws.cell(row=row, column=2, value=text).font = F_BOLD
        for c in range(2, 2 + cols):
            ws.cell(row=row, column=c).fill = FILL_SEC

    # ============================================================ CENNIK
    cen = wb.active
    cen.title = "Cennik"
    header(cen, "CENNIK I STAŁE KALKULACYJNE NK38", 5)
    cen["B3"] = "Zmiana ceny tutaj działa we wszystkich formułach. Kolumna E = skąd wartość pochodzi w starym arkuszu."
    cen["B3"].font = Font(italic=True, color="7F7F7F")
    for col, w in zip("ABCDEF", (2, 54, 14, 14, 36, 12)):
        cen.column_dimensions[col].width = w
    r = 5
    for item in CENNIK:
        if item[1] is None:
            section(cen, r, item[0])
            r += 1
            continue
        n, label, val, fmt, src = item
        cen.cell(row=r, column=2, value=label)
        c = cen.cell(row=r, column=3, value=val)
        c.number_format, c.fill, c.border = fmt, FILL_IN, BOX
        cen.cell(row=r, column=5, value=src).font = F_NOTE
        name(n, cen, f"C{r}")
        r += 1
    r += 1
    section(cen, r, "Zagłębienia (standard 170 cm: bez dopłaty, siłownik 1700)")
    r += 1
    for i, h in enumerate(("Zagłębienie [cm]", "Dopłata", "Siłownik")):
        cen.cell(row=r, column=2 + i, value=h).font = F_BOLD
    r += 1
    g0 = r
    for cm, dop, sil in GLEBOKOSCI:
        cen.cell(row=r, column=2, value=cm)
        cen.cell(row=r, column=3, value=dop).number_format = PLN
        cen.cell(row=r, column=4, value=sil)
        cen.cell(row=r, column=3).fill = cen.cell(row=r, column=4).fill = FILL_IN
        r += 1
    name_range("t_gl_doplata", cen, f"C{g0}:C{r - 1}")
    name_range("t_gl_silownik", cen, f"D{g0}:D{r - 1}")
    r += 1
    section(cen, r, "Składowe platformy (na 1 platformę; N.P. liczone jak D)", 5)
    r += 1
    for i, h in enumerate(("Pozycja", "Cena jedn.", "Ilość na S", "Ilość na D", "Grupa")):
        cen.cell(row=r, column=2 + i, value=h).font = F_BOLD
    r += 1
    s0 = r
    for poz, cena, ils, ild, grp in SKLADOWE:
        for i, v in enumerate((poz, cena, ils, ild, grp)):
            c = cen.cell(row=r, column=2 + i, value=v)
            if i in (1, 2, 3):
                c.fill = FILL_IN
        cen.cell(row=r, column=3).number_format = PLN2
        r += 1
    name_range("t_sk_cena", cen, f"C{s0}:C{r - 1}")
    name_range("t_sk_S", cen, f"D{s0}:D{r - 1}")
    name_range("t_sk_D", cen, f"E{s0}:E{r - 1}")
    name_range("t_sk_grupa", cen, f"F{s0}:F{r - 1}")

    # ============================================================ DANE
    dane = wb.create_sheet("Dane", 0)
    header(dane, "KALKULATOR NK38 - DANE WEJŚCIOWE", 4)
    dane["B3"] = "Wypełniasz tylko żółte pola. Puste pole = 0 lub wartość automatyczna. Wynik: arkusz Podsumowanie."
    dane["B3"].font = Font(italic=True, color="7F7F7F")
    for col, w in zip("ABCDEF", (2, 50, 18, 16, 52, 24)):
        dane.column_dimensions[col].width = w
    dvs = {
        "lok": DataValidation(type="list", formula1='"Kraków,Poza Krakowem"'),
        "typ": DataValidation(type="list", formula1='"NK38,NK37"'),
        "wzor": DataValidation(type="list", formula1='"' + ",".join(WZORY) + '"'),
        "war": DataValidation(type="list", formula1='"' + ",".join(WARIANTY) + '"'),
        "int": DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0"),
        "os": DataValidation(type="whole", operator="between", formula1="1", formula2="12"),
    }
    for dv in dvs.values():
        dane.add_data_validation(dv)

    def inp(row, n, label, default=None, fmt=None, note=None, dv=None, col=3):
        dane.cell(row=row, column=2, value=label)
        c = dane.cell(row=row, column=col, value=inputs.get(n, default) if n else default)
        c.fill, c.border = FILL_IN, BOX
        if fmt:
            c.number_format = fmt
        if note:
            dane.cell(row=row, column=5, value=note).font = F_NOTE
        if dv:
            dvs[dv].add(c)
        if n:
            name(n, dane, f"{get_column_letter(col)}{row}")
        return c

    r = 5
    section(dane, r, "1. Inwestycja"); r += 1
    inp(r, "in_nazwa", "Nazwa / adres inwestycji", "ul. ..."); r += 1
    inp(r, "in_data", "Data oferty", None, "yyyy-mm-dd"); r += 1
    inp(r, "in_nr", "Nr oferty / handlowiec", "SNK"); r += 1
    inp(r, "in_opis", "Opis / wariant (H, L, P, Hpoj ...)", ""); r += 1
    r += 1
    section(dane, r, "2. Lokalizacja"); r += 1
    inp(r, "in_lok", "Lokalizacja", "Poza Krakowem", None,
        "Kraków: montaż w dniach, 20 km, bez hotelu i diet", "lok"); r += 1
    inp(r, "in_km", "Odległość od Krakowa [km, faktyczna, w jedną stronę]", 0, NUM,
        "Pomiary = ta odległość, montaż = +50 km"); r += 1
    inp(r, "in_km_pom", "Nadpisz: km na pomiary / odbiory", None, NUM, "puste = automatycznie"); r += 1
    inp(r, "in_km_mon", "Nadpisz: km na montaż", None, NUM, "puste = automatycznie"); r += 1
    r += 1
    section(dane, r, "3. Urządzenia"); r += 1
    inp(r, "in_typ", "Typ urządzenia", "NK38", None, None, "typ"); r += 1
    inp(r, "in_S", "Platformy pojedyncze S (2 MP)", 0, "0", None, "int"); r += 1
    inp(r, "in_D", "Platformy podwójne D (4 MP)", 0, "0", None, "int"); r += 1
    inp(r, "in_NP", "Systemy dla N.P. (2 MP, liczone jak D)", 0, "0", None, "int"); r += 1
    inp(r, "in_agregaty", "Agregaty", 0, "0", "Zależy od układu garażu (1 agregat do 6 platform)", "int"); r += 1
    inp(r, "in_magistrala", "Magistrala EXTRA [odcinki 6 m]", 0, "0", None, "int"); r += 1
    inp(r, "in_barierki", "Barierki [szt. na stronę]", 0, "0", None, "int"); r += 1
    inp(r, "in_maskownice", "Maskownice / bariery tylne", 0, "0", None, "int"); r += 1
    inp(r, "in_utr_dod", "Utrudnienia - dodatkowe osobodni (ręcznie)", 0, NUM,
        "Doliczane do utrudnień z magistrali i barierek (dawny 'tonaż')"); r += 1
    r += 1
    section(dane, r, "4. Wymiary niestandardowe (ile platform danego typu)"); r += 1
    dane.cell(row=r, column=3, value="S").font = F_BOLD
    dane.cell(row=r, column=4, value="D").font = F_BOLD
    dane.cell(row=r, column=5, value="Pozostałe platformy = wymiar standardowy (250/500, 170, 530)").font = F_NOTE
    r += 1
    gl_first = r
    txt_cells = []
    for cm, _, _ in GLEBOKOSCI:
        inp(r, None, f"Zagłębienie {cm} cm", inputs.get("gl_S", {}).get(cm, 0), "0", None, "int", col=3)
        inp(r, None, f"Zagłębienie {cm} cm", inputs.get("gl_D", {}).get(cm, 0), "0", None, "int", col=4)
        dane.cell(row=r, column=6, value=f'=IF(C{r}+D{r}>0,"; {cm} cm: "&IF(C{r}>0,C{r}&"S","")'
                  f'&IF(AND(C{r}>0,D{r}>0),"+","")&IF(D{r}>0,D{r}&"D",""),"")').font = F_NOTE
        txt_cells.append(f"F{r}")
        r += 1
    name_range("in_gl_S", dane, f"C{gl_first}:C{r - 1}")
    name_range("in_gl_D", dane, f"D{gl_first}:D{r - 1}")
    dane.cell(row=r, column=2, value="Zagłębienie 170 cm (standard) - wyliczone")
    for col, rng, tot in ((3, "in_gl_S", "in_S"), (4, "in_gl_D", "in_D")):
        dane.cell(row=r, column=col, value=f"={tot}-SUM({rng})").fill = FILL_CALC
    r_gl_std = r
    r += 1
    dim_cells = []
    for key, label in WYMIARY:
        inp(r, f"in_{key}_S", label, 0, "0", None, "int", col=3)
        inp(r, f"in_{key}_D", label, 0, "0", None, "int", col=4)
        short = label.replace("Długość zagłębienia ", "dł. ").replace("Szerokość ", "szer. ")
        dane.cell(row=r, column=6, value=f'=IF(C{r}+D{r}>0,"; {short}: "&IF(C{r}>0,C{r}&"S","")'
                  f'&IF(AND(C{r}>0,D{r}>0),"+","")&IF(D{r}>0,D{r}&"D",""),"")').font = F_NOTE
        dim_cells.append(f"F{r}")
        r += 1
    dane.cell(row=r, column=6, value='=MID(' + "&".join(txt_cells) + ',3,500)')
    name("t_zaglebienia", dane, f"F{r}")
    dane.cell(row=r + 1, column=6, value='=MID(' + "&".join(dim_cells) + ',3,500)')
    name("t_wymiary", dane, f"F{r + 1}")
    dane.cell(row=gl_first - 1, column=6, value="(pomocnicze do opisu)").font = F_NOTE
    r += 1
    section(dane, r, "5. Transport i ekipa"); r += 1
    inp(r, "in_tir", "Liczba transportów (TIR)", 0, "0", None, "int"); r += 1
    inp(r, "in_tir_cena", "Cena transportu / TIR", 0, PLN, "Orientacyjnie: KRK 1600, WAW/WRO 3300, POZ/ZG 4500"); r += 1
    inp(r, "in_rozl_cena", "Rozładunek - sprzęt (jednorazowo)", 0, PLN, "Jednorazowy koszt przyjazdu i wyjazdu wózka"); r += 1
    inp(r, "in_rozl_dni", "Nadpisz: dni rozładunku", None, NUM, "puste = 0,5 dnia / TIR"); r += 1
    inp(r, "in_osoby", "Liczba osób na montażu", 4, "0",
        "Stawki: 500 / 400 / każda kolejna 350 zł. Zobacz analizę liczby osób w Podsumowaniu", "os"); r += 1
    inp(r, "in_wyjazdy", "Nadpisz: tygodnie (Kraków: dni) montażu", None, "0", "puste = automatycznie"); r += 1
    inp(r, "in_wzor", "Wzór czasu podróży", WZORY[0], None, "Patrz uwagi w Podsumowaniu", "wzor"); r += 1
    r += 1
    section(dane, r, "6. Narzuty i kurs"); r += 1
    inp(r, "in_zysk", "Zysk A (udział w cenie)", 0.14, PCT); r += 1
    inp(r, "in_zysk2", "Zysk B - do porównania, do ilu możemy zejść", 0.10, PCT); r += 1
    inp(r, "in_q2500", "Dopłata za Q=2500 kg", "=c_q2500", PCT, "domyślnie 7% z Cennika"); r += 1
    inp(r, "in_eur", "Kurs EUR/PLN", 4.325, "0.0000"); r += 1
    inp(r, "in_konsA", "Konserwacja - cena A [zł / MP / miesiąc]", 40, PLN); r += 1
    inp(r, "in_konsB", "Konserwacja - cena B [zł / MP / miesiąc]", 50, PLN); r += 1
    inp(r, "in_analiza", "Analiza liczby osób dla wariantu", WARIANTY[0], None, None, "war"); r += 1
    r += 1
    section(dane, r, "Kontrola danych"); r += 1
    r_warn = r
    warns = [
        f'=IF(OR(C{r_gl_std}<0,D{r_gl_std}<0),"BŁĄD: więcej zagłębień niż platform","ok")',
        '=IF(OR(in_szer1_S+in_szer2_S>in_S,in_szer1_D+in_szer2_D>in_D),"BŁĄD: więcej szerokości niż platform","ok")',
        '=IF(OR(in_dl500_S+in_dl540_S+in_dl550_S>in_S,in_dl500_D+in_dl540_D+in_dl550_D>in_D),"BŁĄD: więcej długości niż platform","ok")',
        '=IF(in_agregaty*6<in_S+in_D+in_NP,"UWAGA: mniej niż 1 agregat na 6 platform","ok")',
        '=IF(in_S+in_D+in_NP=0,"UWAGA: brak platform",IF(in_tir=0,"UWAGA: brak transportów","ok"))',
        '=IF(AND(in_lok="Poza Krakowem",in_km=0),"UWAGA: brak odległości od Krakowa","ok")',
    ]
    for f in warns:
        dane.cell(row=r, column=2, value=f)
        dane.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        r += 1
    r_warn_end = r - 1
    dane.conditional_formatting.add(f"B{r_warn}:B{r_warn_end}",
                                    FormulaRule(formula=[f'B{r_warn}<>"ok"'], fill=FILL_WARN))

    # ============================================================ OBLICZENIA
    obl = wb.create_sheet("Obliczenia")
    header(obl, "OBLICZENIA (nie edytować)", 13)
    obl.column_dimensions["A"].width = 2
    obl.column_dimensions["B"].width = 52
    for i in range(3, 15):
        obl.column_dimensions[get_column_letter(i)].width = 14
    obl.column_dimensions["H"].width = 3

    def calc(row, n, label, f, fmt=NUM, bold=False):
        obl.cell(row=row, column=2, value=label)
        c = obl.cell(row=row, column=3, value=f)
        c.number_format, c.fill = fmt, FILL_CALC
        if bold:
            c.font = obl.cell(row=row, column=2).font = F_BOLD
        name(n, obl, f"C{row}")

    trav = ('IF(in_wzor="Uproszczony",2*{km}/c_km_dzien*{n},'
            '(1+0.1*MAX(0,800-{km})/100)*{km}*{n}/500)')
    grp = 'SUMPRODUCT((t_sk_grupa="{g}")*t_sk_cena*(t_sk_S*in_S+t_sk_D*m_DN))'
    r = 4
    section(obl, r, "A. Materiały i urządzenia (dawny arkusz 'NK38 ...')", 2); r += 1
    MAT = [
        ("m_cena_D", "Cena platformy D", "=c_cena_D+IF(in_typ=\"NK37\",c_doplata_NK37,0)", PLN),
        ("m_cena_S", "Cena platformy S", "=c_wsp_S*m_cena_D", PLN),
        ("m_cena_NP", "Cena systemu N.P.", "=c_wsp_NP*m_cena_D", PLN),
        ("m_DN", "Platformy D + N.P. (N.P. jak D)", "=in_D+in_NP", "0"),
        ("m_platformy", "Razem platform", "=in_S+in_D+in_NP", "0"),
        ("m_MP", "Miejsca postojowe", "=2*in_S+2*in_NP+4*in_D", "0"),
        ("m_plat_S", "Platformy S z dopłatami",
         "=m_cena_S*in_S+c_szer_260*in_szer1_S+c_szer_270*in_szer2_S+SUMPRODUCT(in_gl_S,t_gl_doplata)"
         "+c_dl_500*in_dl500_S+CEILING(m_cena_S*c_dl_wsp*10,100)*in_dl540_S"
         "+CEILING(m_cena_S*c_dl_wsp*20,100)*in_dl550_S", PLN),
        ("m_plat_D", "Platformy D z dopłatami",
         "=m_cena_D*in_D+c_szer_520*in_szer1_D+c_szer_540*in_szer2_D+SUMPRODUCT(in_gl_D,t_gl_doplata)"
         "+c_dl_500*in_dl500_D+m_cena_D*c_dl_wsp*10*in_dl540_D+m_cena_D*c_dl_wsp*20*in_dl550_D", PLN),
        ("m_plat_NP", "Platformy N.P.", "=m_cena_NP*in_NP", PLN),
        ("m_sil_2000", "Siłowniki 2000 [pary]", "=SUMPRODUCT((t_gl_silownik=2000)*(in_gl_S+in_gl_D))", "0"),
        ("m_sil_2250", "Siłowniki 2250 [pary]", "=SUMPRODUCT((t_gl_silownik=2250)*(in_gl_S+in_gl_D))", "0"),
        ("m_sil_std", "Siłowniki standard [pary] (z N.P.)", "=m_platformy-m_sil_2000-m_sil_2250", "0"),
        ("m_hydraulika", "Hydraulika",
         "=in_agregaty*(c_agregat+c_hydr_drobne+c_olej+c_transp_agregatu)+m_platformy*c_transp_silownikow"
         "+2*(c_sil_1700*m_sil_std+c_sil_2000*m_sil_2000+c_sil_2250*m_sil_2250)+in_magistrala*c_magistrala+"
         + grp.format(g="Hydraulika"), PLN),
        ("m_barierki", "Barierki i maskownice", "=in_barierki*c_barierka+in_maskownice*c_maskownica", PLN),
        ("m_elektryka", "Elektryka", "=in_agregaty*c_skrzynka_zasil+" + grp.format(g="Elektryka"), PLN),
        ("m_mechanika", "Mechanika zakupowa i pozostałe", "=in_agregaty*c_admin+" + grp.format(g="Mechanika"), PLN),
        ("m_suma_2000", "SUMA STANDARD Q=2000 (dawne F87)",
         "=m_plat_S+m_plat_D+m_plat_NP+m_hydraulika+m_barierki+m_elektryka+m_mechanika", PLN),
        ("m_suma_2500", "SUMA Q=2500 (dawne F90)", "=m_suma_2000*(1+in_q2500)", PLN),
        ("m_malowanie", "Opcja: malowanie proszkowe RAL (koszt)",
         "=c_mal_S*in_S+c_mal_D*m_DN+c_mal_bar*in_barierki", PLN),
        ("m_utr_auto", "Utrudnienia z magistrali i barierek [osobodni]",
         "=ROUNDUP(in_magistrala*c_osd_magistrala+(in_barierki+in_maskownice)*c_osd_barierka,0)", NUM),
        ("m_utr", "Utrudnienia razem [osobodni] (dawny 'tonaż')", "=m_utr_auto+in_utr_dod", NUM),
    ]
    for n, label, f, fmt in MAT:
        calc(r, n, label, f, fmt, bold=n.startswith("m_suma"))
        r += 1
    r += 1
    section(obl, r, "B. Parametry montażu wspólne dla wariantów", 2); r += 1
    for n, label, f, fmt in [
        ("e_krk", "Kraków?", '=in_lok="Kraków"', "General"),
        ("e_km_pom", "Km na pomiary / odbiory", "=IF(ISBLANK(in_km_pom),IF(e_krk,c_km_krakow,in_km),in_km_pom)", NUM),
        ("e_km_mon", "Km na montaż", "=IF(ISBLANK(in_km_mon),IF(e_krk,c_km_krakow,in_km+c_km_dodatek),in_km_mon)", NUM),
        ("e_rozl_dni", "Dni rozładunku", "=IF(ISBLANK(in_rozl_dni),c_rozl_TIR*in_tir,in_rozl_dni)", NUM),
        ("e_podr_pom", "Podróże pomiary + odbiory [osobodni]",
         "=(c_pomiary_osoby+c_odbiory_osoby)*" + trav.format(km="e_km_pom", n=1), NUM),
        ("e_dzw_dni", "Dni dźwigu", "=ROUNDUP(m_DN/c_dzwig_D+in_S*c_dzwig_wspS/c_dzwig_D,0)", "0"),
        ("e_mnoznik", "Mnożnik kosztu materiałowego na cenę (opcje)",
         "=(1+c_podatek*(c_gdw+in_zysk))/(1-(c_kob+c_gdw+in_zysk))", "0.0000"),
    ]:
        calc(r, n, label, f, fmt)
        r += 1
    r += 1

    # --- montaż w kolumnach: 4 warianty (D:G) + analiza liczby osób (I:N)
    section(obl, r, "C. Montaż i cena (dawny arkusz '<Inwestycja>-NK38')", 13); r += 1
    cols = []
    for i, w in enumerate(WARIANTY):
        cols.append(dict(col=get_column_letter(4 + i), label=w,
                         q=1 if w.startswith("2500") else 0, sl=1 if "słupki" in w else 0,
                         os="=in_osoby"))
    for i, n in enumerate(OSOBY_ANALIZA):
        cols.append(dict(col=get_column_letter(9 + i), label=f"{n} os.",
                         q='=IF(LEFT(in_analiza,4)="2500",1,0)',
                         sl='=IF(ISNUMBER(SEARCH("słupki",in_analiza)),1,0)', os=n))
    obl.cell(row=r, column=2, value="Wariant / liczba osób").font = F_BOLD
    obl.cell(row=r - 1, column=9, value="Analiza liczby osób (wariant z Dane)").font = F_BOLD
    for c in cols:
        cell = obl[f"{c['col']}{r}"]
        cell.value, cell.font = c["label"], F_BOLD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    rows = {}

    def mrow(key, label, tmpl, fmt=PLN, bold=False):
        nonlocal r
        obl.cell(row=r, column=2, value=label)
        rows[key] = r
        for c in cols:
            f = tmpl(c) if callable(tmpl) else tmpl
            if isinstance(f, str):
                for k, v in rows.items():
                    f = f.replace("{" + k + "}", f"{c['col']}{v}")
            cell = obl[f"{c['col']}{r}"]
            cell.value, cell.number_format, cell.fill = f, fmt, FILL_CALC
            if bold:
                cell.font = F_BOLD
        if bold:
            obl.cell(row=r, column=2).font = F_BOLD
        r += 1

    def podr(t):
        return "(e_podr_pom+{os}*(e_rozl_dni+" + trav.format(km="e_km_mon", n=t) + "))"

    mrow("q", "Q=2500? (1/0)", lambda c: c["q"], "0")
    mrow("sl", "Słupki e-ładowarek? (1/0)", lambda c: c["sl"], "0")
    mrow("os", "Liczba osób", lambda c: c["os"], "0")
    mrow("st", "Średnia stawka dzienna",
         "=(c_stawka_1+IF({os}>=2,c_stawka_2,0)+MAX(0,{os}-2)*c_stawka_n)/{os}")
    mrow("osd", "Osobodni montażu (dawne F50)",
         "=c_osd_S*in_S+c_osd_D*in_D+c_osd_NP*in_NP+m_utr+{sl}*c_slupek_rbd*m_MP", NUM)
    prev = "1"
    for it in range(1, 6):
        mrow(f"it{it}", f"Iteracja {it}: tygodnie / dni montażu",
             "=IF(e_krk,{d},CEILING({d}/c_dni_tydzien,1))".replace(
                 "{d}", "CEILING(({osd}+" + podr(prev) + ")/{os},1)"), "0")
        prev = "{it%d}" % it
    mrow("wyj", "Tygodnie (Kraków: dni) montażu (dawne B55)",
         "=IF(ISBLANK(in_wyjazdy),{it5},in_wyjazdy)", "0", True)
    mrow("podr", "Osobodni podróży i rozładunku (dawne G59)", "=" + podr("{wyj}"), NUM)
    mrow("osd_all", "Osobodni razem", "={osd}+{podr}", NUM)
    mrow("dni", "Dni pracy ekipy (dawne J39)", "=CEILING({osd_all}/{os},1)", "0", True)
    mrow("tyg", "Tygodnie montażu (do oferty)", "=IF(e_krk,CEILING({dni}/c_dni_tydzien,1),{wyj})", "0", True)
    mrow("km", "Kilometry razem (1 auto ekipy)", "=2*(e_km_mon*{wyj}+2*e_km_pom)", NUM)
    mrow("k_urz", "Urządzenia (materiały)", "=IF({q}=1,m_suma_2500,m_suma_2000)")
    mrow("k_sl", "Słupki e-ładowarek", "={sl}*c_slupek*m_MP")
    mrow("k_pal", "Paliwo", "={km}/100*ROUNDUP(c_spalanie*c_paliwo,0)")
    mrow("k_hot", "Hotel", "=IF(e_krk,0,c_hotel*({dni}+1)*{os})")
    mrow("k_tr", "Transport TIR", "=in_tir_cena*in_tir")
    mrow("k_dzw", "Spider Crane", "=c_dzwig*e_dzw_dni")
    mrow("k_rs", "Rozładunek (sprzęt, jednorazowo)", "=in_rozl_cena")
    mrow("k_rl", "Rozładunek (ludzie)", "={st}*e_rozl_dni*{os}")
    mrow("k_doj", "Dojazdy (ludzie)", "={st}*({podr}-{os}*e_rozl_dni)")
    mrow("k_mon", "Montaż", "=IF({wyj}>0,{st}*{osd},0)")
    mrow("k_die", "Diety", "=IF(e_krk,0,c_dieta*{osd_all})")
    mrow("baza", "Suma kosztów przed narzutami", "=SUM({k_urz}:{k_die})", PLN, True)
    mrow("brutto", "Podstawa narzutów", "={baza}/(1-(c_kob+c_gdw+in_zysk))")
    mrow("k_kob", "KOB", "={brutto}*c_kob")
    mrow("k_gdw", "GDW", "={brutto}*c_gdw")
    mrow("k_zysk", "Zysk", "={brutto}*in_zysk")
    mrow("k_pod", "Podatek", "=c_podatek*({k_rl}+{k_doj}+{k_mon}+{k_die}+{k_gdw}+{k_zysk})")
    mrow("cena", "CENA KONTRAKTU NETTO (dawne K39)", "={brutto}+{k_pod}", PLN, True)
    mrow("koszt", "Koszt rzeczywisty (dawne K43)", "={cena}-{k_zysk}")
    mrow("cena_mp", "Cena za MP", "=IF(m_MP>0,{cena}/m_MP,0)")
    mrow("cena_mp_eur", "Cena za MP [EUR]", "={cena_mp}/in_eur", EUR)
    mrow("brutto2", "Zysk B: podstawa narzutów", "={baza}/(1-(c_kob+c_gdw+in_zysk2))")
    mrow("k_zysk2", "Zysk B", "={brutto2}*in_zysk2")
    mrow("cena2", "CENA KONTRAKTU NETTO - ZYSK B",
         "={brutto2}+c_podatek*({k_rl}+{k_doj}+{k_mon}+{k_die}+{brutto2}*(c_gdw+in_zysk2))", PLN, True)
    mrow("cena_mp2", "Cena za MP - zysk B", "=IF(m_MP>0,{cena2}/m_MP,0)")

    def oref(key, i):
        return f"Obliczenia!{cols[i]['col']}{rows[key]}"

    # ============================================================ PODSUMOWANIE
    pod = wb.create_sheet("Podsumowanie", 1)
    header(pod, "PODSUMOWANIE OFERTY NK38", 7)
    for col, w in zip("ABCDEFGHI", (2, 44, 17, 17, 17, 17, 17, 17, 2)):
        pod.column_dimensions[col].width = w
    pod["B3"] = '=in_nazwa&"   |   "&IF(ISBLANK(in_data),"",TEXT(in_data,"yyyy-mm-dd"))&"   |   "&in_nr'
    pod["B3"].font = F_BOLD
    pod["B4"] = ('=in_typ&" "&IF(in_S>0,in_S&"S250","")&IF(AND(in_S>0,in_D>0),"+","")'
                 '&IF(in_D>0,in_D&"D500","")&IF(in_NP>0,"+"&in_NP&"NP","")'
                 '&", "&in_agregaty&"A   |   "&m_MP&" MP   |   "&in_lok')
    r = 6

    def table_head(row, labels, first="Wariant"):
        pod.cell(row=row, column=2, value=first).font = F_BOLD
        pod.cell(row=row, column=2).fill = FILL_SEC
        for i, lab in enumerate(labels):
            c = pod.cell(row=row, column=3 + i, value=lab)
            c.font, c.fill = F_BOLD, FILL_SEC
            c.alignment = Alignment(horizontal="center")

    table_head(r, WARIANTY); r += 1
    MAIN = [
        ('="CENA NETTO - ZYSK A "&TEXT(in_zysk,"0.0%")', ("cena",), PLN, True),
        ('="CENA NETTO - ZYSK B "&TEXT(in_zysk2,"0.0%")', ("cena2",), PLN, True),
        ("Cena za MP - zysk A", ("cena_mp",), PLN, False),
        ("Cena za MP - zysk B", ("cena_mp2",), PLN, False),
        ("Cena za MP [EUR] - zysk A", ("cena_mp_eur",), EUR, False),
        ("Zysk A [zł]", ("k_zysk",), PLN, False),
        ("Zysk B [zł]", ("k_zysk2",), PLN, False),
        ("Koszt rzeczywisty", ("koszt",), PLN, False),
        (None,) * 4,
        ("Urządzenia (materiały)", ("k_urz",), PLN, False),
        ("Słupki e-ładowarek", ("k_sl",), PLN, False),
        ("Transport, dźwig, rozładunek", ("k_tr", "k_dzw", "k_rs"), PLN, False),
        ("Robocizna (montaż, rozładunek, dojazdy)", ("k_mon", "k_rl", "k_doj"), PLN, False),
        ("Hotel, diety, paliwo", ("k_hot", "k_die", "k_pal"), PLN, False),
        ("KOB + GDW + podatek", ("k_kob", "k_gdw", "k_pod"), PLN, False),
        (None,) * 4,
        ("Dni pracy ekipy", ("dni",), "0", True),
        ("Tygodnie montażu", ("tyg",), "0", True),
        ("Osobodni razem", ("osd_all",), NUM, False),
    ]
    for label, keys, fmt, bold in MAIN:
        if label is None:
            r += 1
            continue
        pod.cell(row=r, column=2, value=label)
        for i in range(4):
            c = pod.cell(row=r, column=3 + i, value="=" + "+".join(oref(k, i) for k in keys))
            c.number_format, c.border = fmt, BOX
            if bold:
                c.font, c.fill = F_BOLD, FILL_RES
        if bold:
            pod.cell(row=r, column=2).font = F_BOLD
        r += 1

    r += 1
    section(pod, r, "Opcje do oferty (cena z narzutami)", 3)
    pod.cell(row=r, column=3, value="Q=2000 kg").font = F_BOLD
    pod.cell(row=r, column=4, value="Q=2500 kg").font = F_BOLD
    r += 1
    for label, f2000, f2500 in [
        ("Słupki e-ładowarek 3,7 kW (z robocizną)", f"={oref('cena', 2)}-{oref('cena', 0)}",
         f"={oref('cena', 3)}-{oref('cena', 1)}"),
        ("Słupki e-ładowarek 22 kW (z robocizną)",
         f"={oref('cena', 2)}-{oref('cena', 0)}+(c_slupek22-c_slupek)*m_MP*e_mnoznik",
         f"={oref('cena', 3)}-{oref('cena', 1)}+(c_slupek22-c_slupek)*m_MP*e_mnoznik"),
        ("Malowanie proszkowe RAL", "=m_malowanie*e_mnoznik", "=m_malowanie*e_mnoznik"),
    ]:
        pod.cell(row=r, column=2, value=label)
        for i, f in enumerate((f2000, f2500)):
            c = pod.cell(row=r, column=3 + i, value=f)
            c.number_format, c.border = PLN, BOX
        r += 1

    r += 1
    pod.cell(row=r, column=2, value='="Analiza liczby osób - wariant: "&in_analiza').font = F_BOLD
    r += 1
    table_head(r, [f"{n} os." for n in OSOBY_ANALIZA], "Liczba osób"); r += 1
    r_best = None
    for label, key, fmt in [("Cena kontraktu", "cena", PLN), ("Cena za MP", "cena_mp", PLN),
                            ("Dni pracy ekipy", "dni", "0"), ("Tygodnie (Kraków: dni)", "wyj", "0"),
                            ("Hotel + diety", None, PLN)]:
        pod.cell(row=r, column=2, value=label)
        for j in range(len(OSOBY_ANALIZA)):
            i = 4 + j
            f = f"={oref(key, i)}" if key else f"={oref('k_hot', i)}+{oref('k_die', i)}"
            c = pod.cell(row=r, column=3 + j, value=f)
            c.number_format, c.border = fmt, BOX
        if key == "cena":
            r_best = r
        r += 1
    last = get_column_letter(2 + len(OSOBY_ANALIZA))
    pod.conditional_formatting.add(f"C{r_best}:{last}{r_best}",
                                   FormulaRule(formula=[f"C{r_best}=MIN($C${r_best}:${last}${r_best})"],
                                               fill=FILL_BEST))
    pod.cell(row=r, column=2, value="Zielone = najtańsza liczba osób").font = F_NOTE
    r += 2

    section(pod, r, "Zestawienie do oferty", 5); r += 1
    ZEST = [
        ("Urządzenia", '=in_typ&": "&IF(in_S>0,in_S&" x S (2 MP)  ","")&IF(in_D>0,in_D&" x D (4 MP)  ","")'
                       '&IF(in_NP>0,in_NP&" x N.P.  ","")&"= "&m_MP&" MP"'),
        ("Agregaty", '=in_agregaty&" szt."'),
        ("Zagłębienia niestandardowe", '=IF(t_zaglebienia="","standard 170 cm",t_zaglebienia)'),
        ("Wymiary niestandardowe", '=IF(t_wymiary="","standard",t_wymiary)'),
        ("Magistrala EXTRA / barierki / maskownice",
         '=in_magistrala&" x 6 m  /  "&in_barierki&" szt.  /  "&in_maskownice&" szt."'),
        ("Utrudnienia [osobodni]", '=m_utr&" (z magistrali i barierek "&m_utr_auto&", ręcznie "&in_utr_dod&")"'),
        ("Lokalizacja", '=in_lok&IF(e_krk,"","  ("&in_km&" km od Krakowa)")'),
        ("Km pomiary / montaż", '=e_km_pom&" / "&e_km_mon&" km"'),
        ("Transport", '=in_tir&" TIR x "&TEXT(in_tir_cena,"#,##0")&" zł, rozładunek sprzętu "&TEXT(in_rozl_cena,"#,##0")&" zł jednorazowo, "&e_rozl_dni&" dni"'),
        ("Ekipa", f'=in_osoby&" osób, średnia stawka "&TEXT({oref("st", 0)},"#,##0")&" zł/dzień"'),
        ("Dźwig", '=e_dzw_dni&" dni"'),
        ("Czas montażu (bez słupków)",
         f'={oref("dni", 0)}&" dni pracy = "&{oref("wyj", 0)}&IF(e_krk," dni"," tyg.")&", "&TEXT({oref("osd_all", 0)},"0.0")&" osobodni"'),
        ("Czas montażu (ze słupkami)",
         f'={oref("dni", 2)}&" dni pracy = "&{oref("wyj", 2)}&IF(e_krk," dni"," tyg.")&", "&TEXT({oref("osd_all", 2)},"0.0")&" osobodni"'),
        ("Zysk A / zysk B", '=TEXT(in_zysk,"0.0%")&" / "&TEXT(in_zysk2,"0.0%")'),
        ("Dopłata Q=2500", '=TEXT(in_q2500,"0.0%")'),
        ("Wzór czasu podróży", "=in_wzor"),
        ("Konserwacja roczna A / B",
         '=TEXT(in_konsA*m_MP*12,"#,##0")&" / "&TEXT(in_konsB*m_MP*12,"#,##0")&" zł  (przegląd półroczny "&TEXT(in_konsA*m_MP*6,"#,##0")&" / "&TEXT(in_konsB*m_MP*6,"#,##0")&" zł)"'),
        ("Opis / wariant", '=in_opis&""'),
    ]
    for label, f in ZEST:
        pod.cell(row=r, column=2, value=label)
        pod.cell(row=r, column=3, value=f)
        pod.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        r += 1
    r += 1
    pod.cell(row=r, column=2, value="Kontrola danych:").font = F_BOLD
    n_w = r_warn_end - r_warn + 1
    pod.cell(row=r, column=3, value=f'=IF(COUNTIF(Dane!B{r_warn}:B{r_warn_end},"ok")={n_w},"wszystko ok","sprawdź arkusz Dane")')
    pod.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f'C{r}<>"wszystko ok"'], fill=FILL_WARN))
    r += 2
    section(pod, r, "Uwagi do sprawdzenia", 7); r += 1
    for u in UWAGI:
        pod.cell(row=r, column=2, value="• " + u).alignment = Alignment(wrap_text=True, vertical="top")
        pod.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        pod.row_dimensions[r].height = 30
        r += 1

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.active = 0
    return wb, names, rows, cols


if __name__ == "__main__":
    args = sys.argv[1:]
    out = args[0] if args and not args[0].startswith("--") else "Kalkulator_NK38_v1.xlsx"
    inputs = PRZYKLADY[args[args.index("--przyklad") + 1]] if "--przyklad" in args else {}
    build(inputs)[0].save(out)
    print("zapisano", out)
