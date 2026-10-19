"""Całościowe opracowanie dziejów rodziny: składanie tekstu i plik Word.

Użycie: python narzedzia/historia.py
Wejście: teksty/historia/*.md (części w kolejności nazw) — Markdown z przypisami [^etykieta] i odsyłaczami ⟨S..-p..-..⟩.
Wynik:
  teksty/historia-rodziny.md                         — tekst złożony (z niego buduj.py robi stronę historia.html)
  Opracowanie/Historia rodziny Boxa-Radoszewskich.docx — wersja do druku: przypisy dolne, spis treści, ilustracje
Ilustracje: ![podpis](skany/plik.jpg) — pliki z strona/skany/ (narzedzia/zdjecia_www.py).
"""
import datetime, glob, html, os, re, sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KORZEN = os.path.normpath(os.path.join(os.path.dirname(__file__), '..'))
CZESCI = os.path.join(KORZEN, 'teksty', 'historia')
WYJ_MD = os.path.join(KORZEN, 'teksty', 'historia-rodziny.md')
WYJ_DIR = os.path.join(KORZEN, 'Opracowanie')
WYJ_DOCX = os.path.join(WYJ_DIR, 'Historia rodziny Boxa-Radoszewskich.docx')
SKANY = os.path.join(KORZEN, 'strona', 'skany')

AKCENT = RGBColor(0x7A, 0x1F, 0x1F)
SZARY = RGBColor(0x6B, 0x62, 0x58)
FONT = 'Cambria'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


# ============================================================ Markdown → bloki

def bloki(tekst):
    """Prosty parser używanego podzbioru Markdown. Zwraca (bloki, przypisy)."""
    przyp, wyn, akapit, tab, cyt = {}, [], [], [], []

    def zamknij():
        nonlocal akapit, tab, cyt
        if akapit:
            wyn.append(('p', ' '.join(akapit)))
            akapit = []
        if tab:
            wyn.append(('tab', tab))
            tab = []
        if cyt:
            wyn.append(('cyt', ' '.join(cyt)))
            cyt = []

    for w in tekst.split('\n'):
        m = re.match(r'^\[\^([^\]]+)\]:\s*(.*)$', w)
        if m:
            zamknij()
            przyp[m.group(1)] = m.group(2)
            continue
        s = w.strip()
        if not s:
            zamknij()
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            zamknij()
            wyn.append((f'h{len(m.group(1)) - 1}', m.group(2)))
            continue
        if s == '---':
            zamknij()
            wyn.append(('hr', ''))
            continue
        m = re.match(r'^!\[(.*)\]\((.*)\)$', s)
        if m:
            zamknij()
            wyn.append(('img', m.group(1), m.group(2)))
            continue
        if s.startswith('|'):
            if akapit or cyt:
                zamknij()
            if not re.match(r'^\|[\s:|-]+\|$', s):
                tab.append([c.strip() for c in s.strip('|').split('|')])
            continue
        if s.startswith('>'):
            if akapit or tab:
                zamknij()
            cyt.append(s.lstrip('>').strip())
            continue
        m = re.match(r'^(\s*)(?:([-*+])|(\d+)\.)\s+(.*)$', w)
        if m:
            zamknij()
            wyn.append(('li', len(m.group(1).replace('\t', '  ')) // 2, bool(m.group(3)), m.group(4)))
            continue
        if wyn and wyn[-1][0] == 'li' and w.startswith('  ') and not akapit:
            t, poz, num, tx = wyn[-1]
            wyn[-1] = (t, poz, num, tx + ' ' + s)
            continue
        akapit.append(s)
    zamknij()
    return wyn, przyp


INLINE = re.compile(r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|(?<![\w*])\*[^*\n]+?\*(?![\w*])|`[^`]+`|\[\^[^\]]+\]|⟨[^⟩]*⟩)')


def kawalki(tekst):
    """Tekst z formatowaniem → [(tekst, {'b','i','code','przyp','seg'})]."""
    wyn = []
    for k in INLINE.split(tekst):
        if not k:
            continue
        if k.startswith('***') and k.endswith('***') and len(k) > 6:
            wyn += [(t, f | {'b', 'i'}) for t, f in kawalki(k[3:-3])]
        elif k.startswith('**') and k.endswith('**') and len(k) > 4:
            wyn += [(t, f | {'b'}) for t, f in kawalki(k[2:-2])]
        elif k.startswith('*') and k.endswith('*') and len(k) > 2 and not k.startswith('**'):
            wyn += [(t, f | {'i'}) for t, f in kawalki(k[1:-1])]
        elif k.startswith('`') and k.endswith('`'):
            wyn.append((k[1:-1], {'code'}))
        elif k.startswith('[^'):
            wyn.append((k[2:-1], {'przyp'}))
        elif k.startswith('⟨'):
            wyn.append((k[1:-1], {'seg'}))
        else:
            wyn.append((k, set()))
    return wyn


# ============================================================ Word: pomocnicze

def pole(p, instr, tekst=''):
    """Pole Worda (PAGE, TOC …) w akapicie."""
    def fld(t):
        e = OxmlElement('w:fldChar')
        e.set(qn('w:fldCharType'), t)
        r = OxmlElement('w:r')
        r.append(e)
        return r
    p._p.append(fld('begin'))
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r.append(it)
    p._p.append(r)
    p._p.append(fld('separate'))
    if tekst:
        r = p.add_run(tekst)
        r.font.color.rgb = SZARY
    p._p.append(fld('end'))


class Przypisy:
    def __init__(self, defs):
        self.defs, self.lista = defs, []

    def odsylacz(self, p, etykieta):
        self.lista.append(self.defs.get(etykieta, f'[brak przypisu {etykieta}]'))
        n = len(self.lista)
        r = p.add_run()
        r.style = 'Footnote Reference'
        ref = OxmlElement('w:footnoteReference')
        ref.set(qn('w:id'), str(n))
        r._r.append(ref)

    def xml(self):
        def run(t, f):
            rpr = ''
            if 'b' in f:
                rpr += '<w:b/>'
            if 'i' in f:
                rpr += '<w:i/>'
            if 'seg' in f:
                rpr += '<w:color w:val="8A8178"/><w:sz w:val="15"/>'
                t = f'[{t}]'
            if 'code' in f:
                rpr += '<w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>'
            return (f'<w:r>{"<w:rPr>" + rpr + "</w:rPr>" if rpr else ""}'
                    f'<w:t xml:space="preserve">{html.escape(t, quote=False)}</w:t></w:r>')
        sep = ('<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>')
        cz = []
        for n, t in enumerate(self.lista, 1):
            runs = ''.join(run(x, f) for x, f in kawalki(' ' + t))
            cz.append(f'<w:footnote w:id="{n}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
                      f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>{runs}</w:p></w:footnote>')
        return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W_NS}" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">' + sep + ''.join(cz) + '</w:footnotes>').encode('utf-8')


def dodaj_tekst(p, tekst, P, rozmiar=None):
    for t, f in kawalki(tekst):
        if 'przyp' in f:
            P.odsylacz(p, t)
            continue
        if 'seg' in f:
            t = f'[{t}]'
        r = p.add_run(t)
        r.bold = 'b' in f or None
        r.italic = 'i' in f or None
        if 'code' in f:
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        if 'seg' in f:
            r.font.size = Pt(8)
            r.font.color.rgb = SZARY
        if rozmiar:
            r.font.size = rozmiar


def style(doc):
    st = doc.styles
    n = st['Normal']
    n.font.name = FONT
    n.font.size = Pt(11)
    n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    for nazwa, rozm, przed in (('Title', 28, 0), ('Heading 1', 18, 12), ('Heading 2', 14, 14), ('Heading 3', 12, 10)):
        s = st[nazwa]
        s.font.name = FONT
        s.font.size = Pt(rozm)
        s.font.color.rgb = AKCENT
        s.font.bold = nazwa != 'Title'
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rf.set(qn(a), FONT)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        s.paragraph_format.space_before = Pt(przed)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    ft = st.add_style('Footnote Text', WD_STYLE_TYPE.PARAGRAPH)
    ft.base_style = n
    ft.font.size = Pt(8.5)
    ft.paragraph_format.space_after = Pt(1)
    ft.paragraph_format.line_spacing = 1.0
    fr = st.add_style('Footnote Reference', WD_STYLE_TYPE.CHARACTER)
    fr.font.superscript = True
    fr.font.color.rgb = AKCENT
    q = st['Quote']
    q.font.size = Pt(10.5)
    q.paragraph_format.left_indent = Cm(1)
    q.paragraph_format.right_indent = Cm(0.5)
    st['Caption'].font.size = Pt(9)
    st['Caption'].font.color.rgb = SZARY


def okladka(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('⚔')
    r.font.size = Pt(36)
    r.font.color.rgb = AKCENT
    t = doc.add_paragraph('Boxa-Radoszewscy herbu Oksza', style='Title')
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tekst, rozm, kol in (('Dzieje rodziny od legendy o Wrszowcach do czasów współczesnych', 16, AKCENT),
                             ('', 10, None),
                             ('Opracowanie na podstawie rodzinnego zbioru dokumentów, zapisków i publikacji (źródła S01–S10), '
                              'z przypisami do każdej informacji', 11, SZARY),
                             ('', 10, None),
                             ('Wraz z rodzinami spokrewnionymi: Prządzyńskimi h. Grzymała, Drohojowskimi h. Korczak, '
                              'Grabińskimi h. Pomian, Skrzyńskimi h. Zaremba i in.', 11, SZARY)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if tekst:
            r = p.add_run(tekst)
            r.font.size = Pt(rozm)
            if kol:
                r.font.color.rgb = kol
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(160)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mies = ['stycznia', 'lutego', 'marca', 'kwietnia', 'maja', 'czerwca', 'lipca', 'sierpnia', 'września',
            'października', 'listopada', 'grudnia']
    d = datetime.date.today()
    r = p.add_run(f'Wersja robocza do przejrzenia przez rodzinę · {d.day} {mies[d.month - 1]} {d.year}')
    r.font.size = Pt(10)
    r.font.color.rgb = SZARY
    doc.add_page_break()
    h = doc.add_paragraph()
    r = h.add_run('Spis treści')
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), AKCENT
    p = doc.add_paragraph()
    pole(p, 'TOC \\o "1-2" \\h \\z \\u', 'Spis treści uzupełni się po otwarciu pliku (Word zapyta o aktualizację pól; '
                                       'można też kliknąć tu prawym przyciskiem i wybrać „Aktualizuj pole”).')
    doc.add_page_break()


def stopka(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Boxa-Radoszewscy herbu Oksza — ')
    r.font.size = Pt(8.5)
    r.font.color.rgb = SZARY
    pole(p, 'PAGE')


def tabela(doc, wiersze, P):
    kol = max(len(w) for w in wiersze)
    t = doc.add_table(rows=0, cols=kol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(wiersze):
        cells = t.add_row().cells
        for j in range(kol):
            c = cells[j]
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            dodaj_tekst(c.paragraphs[0], w[j] if j < len(w) else '', P, Pt(9))
            if i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True
                tcpr = c._tc.get_or_add_tcPr()
                sh = OxmlElement('w:shd')
                sh.set(qn('w:val'), 'clear')
                sh.set(qn('w:fill'), 'EFE8DC')
                tcpr.append(sh)
    # nagłówek powtarzany na kolejnych stronach
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trpr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def word(blk, przyp):
    doc = Document()
    style(doc)
    stopka(doc)
    okladka(doc)
    P = Przypisy(przyp)
    # okładka zastępuje tytuł, podtytuł i pierwszy akapit z kursywą
    start = next(i for i, b in enumerate(blk) if b[0] == 'h2')
    pierwszy_h1 = True
    for b in blk[start:]:
        typ = b[0]
        if typ == 'h1':
            h = doc.add_paragraph(style='Heading 1')
            dodaj_tekst(h, b[1], P)
            h.paragraph_format.page_break_before = not pierwszy_h1
            pierwszy_h1 = False
        elif typ in ('h2', 'h3'):
            h = doc.add_paragraph(style='Heading 2' if typ == 'h2' else 'Heading 3')
            dodaj_tekst(h, b[1], P)
        elif typ == 'p':
            dodaj_tekst(doc.add_paragraph(), b[1], P)
        elif typ == 'cyt':
            dodaj_tekst(doc.add_paragraph(style='Quote'), b[1], P)
        elif typ == 'li':
            _, poz, num, tx = b
            st = ('List Number' if num else 'List Bullet') + (' 2' if poz else '')
            p = doc.add_paragraph(style=st)
            p.paragraph_format.space_after = Pt(3)
            dodaj_tekst(p, tx, P)
        elif typ == 'tab':
            tabela(doc, b[1], P)
        elif typ == 'img':
            plik = os.path.join(SKANY, os.path.basename(b[2]))
            if os.path.exists(plik):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(plik, width=Cm(13.5))
                c = doc.add_paragraph(style='Caption')
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                dodaj_tekst(c, b[1], P)
            else:
                print('Brak ilustracji:', plik, file=sys.stderr)
        elif typ == 'hr':
            pass
    # przypisy dolne: osobna część dokumentu
    blob = P.xml()
    fn = next((r.target_part for r in doc.part.rels.values() if r.reltype == RT.FOOTNOTES), None)
    if fn is not None:
        fn._blob = blob
    else:
        part = Part(PackURI('/word/footnotes.xml'),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml', blob, doc.part.package)
        doc.part.relate_to(part, RT.FOOTNOTES)
    # przy otwarciu Word zaktualizuje spis treści
    ust = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    ust.append(uf)
    doc.core_properties.title = 'Boxa-Radoszewscy herbu Oksza. Dzieje rodziny'
    doc.core_properties.subject = 'Historia rodziny Radoszewskich h. Oksza na podstawie źródeł S01–S10'
    doc.core_properties.language = 'pl-PL'
    os.makedirs(WYJ_DIR, exist_ok=True)
    doc.save(WYJ_DOCX)
    return len(P.lista)


def main():
    czesci = sorted(glob.glob(os.path.join(CZESCI, '*.md')))
    tekst = '\n\n'.join(open(p, encoding='utf-8').read().strip() for p in czesci) + '\n'
    open(WYJ_MD, 'w', encoding='utf-8').write(
        '<!-- Plik złożony automatycznie z teksty/historia/*.md przez narzedzia/historia.py — nie edytować ręcznie. -->\n\n' + tekst)
    blk, przyp = bloki(tekst)
    uzyte = set(re.findall(r'\[\^([^\]]+)\](?!:)', re.sub(r'^\[\^[^\]]+\]:.*$', '', tekst, flags=re.M)))
    brak = sorted(uzyte - set(przyp))
    nieuz = sorted(set(przyp) - uzyte)
    if brak:
        print('Odsyłacze bez treści przypisu:', ', '.join(brak))
    if nieuz:
        print('Przypisy nieużyte w tekście:', ', '.join(nieuz))
    n = word(blk, przyp)
    slowa = len(re.findall(r'\w+', re.sub(r'^\[\^[^\]]+\]:.*$', '', tekst, flags=re.M)))
    print(f'Złożono {len(czesci)} części: ok. {slowa} słów tekstu głównego, {n} przypisów dolnych.')
    print(f'Zapisano: {os.path.relpath(WYJ_MD, KORZEN)} i {os.path.relpath(WYJ_DOCX, KORZEN)}')


if __name__ == '__main__':
    main()
